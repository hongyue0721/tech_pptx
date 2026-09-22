"""T10 循环②：ExportService 单元测试（先红后绿）。

契约真源 api.md §导出：POST exports 202 JobAccepted；export 是只读快照 job
（不占项目写锁、可排队，job_repository claim_next 已按 kind='export' 放行）；
固定 version 纯渲染；产物经 ArtifactStore 原子落盘（pptx + evidence-report）。
"""

import copy
import io
import json
import zipfile
from datetime import datetime, timezone

import pytest
from pptx import Presentation

from courseware_core.errors import DeckExportError, DeckNotFound, ProjectNotFound
from courseware_core.models import DeckSpec
from courseware_core.models.export import ExportRequest
from courseware_core.render.exporter import EXPORTER_VERSION
from courseware_core.services.export_service import EXPORT_DEADLINE_SECONDS, ExportService
from courseware_core.storage.artifact_store import ArtifactStore
from courseware_core.storage.version_repository import VersionRepository

COURSE = {
    "topic": "STM32 中断机制",
    "audience": "大二学生",
    "duration_minutes": 45,
    "goals": ["理解 NVIC"],
    "target_slides": 8,
}


def deck(version: int, text: str, corpus_revision: int = 1) -> DeckSpec:
    return DeckSpec(
        schema_version="1.0.0",
        project_id="prj_001",
        version=version,
        corpus_revision=corpus_revision,
        course=COURSE,
        claims=[],
        slides=[
            {
                "id": "sld_001",
                "title": f"页 {text}",
                "layout": "title",
                "blocks": [{"type": "teaching", "text": text}],
            }
        ],
    )


@pytest.fixture()
def seeded(conn, insert_project, tmp_path):
    insert_project("prj_001", current_version=0, corpus_revision=1)
    versions = VersionRepository(conn)
    versions.commit_version(
        "prj_001", deck(1, "one"), expected_base_version=0, expected_corpus_revision=1
    )
    versions.commit_version(
        "prj_001", deck(2, "two"), expected_base_version=1, expected_corpus_revision=1
    )
    return conn, tmp_path / "artifacts"


def claim_running(conn, job_id):
    # 模拟 worker 领取语义（queued→running）：handle_export 落盘前复核执行态。
    with conn:
        conn.execute(
            "UPDATE jobs SET status='running', worker_id='w_test' WHERE id=?", (job_id,)
        )
        assert conn.execute("SELECT changes() AS n").fetchone()["n"] == 1


def run_export(svc, conn, job_id):
    claim_running(conn, job_id)
    return svc.handle_export(job_id)


def service(conn, root) -> ExportService:
    return ExportService(conn, artifacts_root=root)


class TestCreateExportJob:
    def test_queues_without_touching_write_lock(self, seeded):
        conn, root = seeded
        job = service(conn, root).create_export_job(
            "prj_001", ExportRequest(version=1, format="pptx")
        )
        assert job.kind == "export"
        assert job.status == "queued"
        assert job.base_version == 1
        assert job.corpus_revision == 1
        # 只读快照：项目写锁指针必须仍为空（export 不排他）。
        row = conn.execute("SELECT active_job_id FROM projects WHERE id='prj_001'").fetchone()
        assert row["active_job_id"] is None

    def test_accepted_even_when_project_busy(self, seeded):
        # 与 claim_next 的 kind='export' 放行同口径：读型 job 不受写锁门拦截。
        conn, root = seeded
        with conn:
            conn.execute(
                "UPDATE projects SET active_job_id='job_other' WHERE id='prj_001'"
            )
        job = service(conn, root).create_export_job(
            "prj_001", ExportRequest(version=2, format="pptx")
        )
        assert job.status == "queued"

    def test_missing_version_raises_deck_not_found(self, seeded):
        conn, root = seeded
        with pytest.raises(DeckNotFound):
            service(conn, root).create_export_job(
                "prj_001", ExportRequest(version=99, format="pptx")
            )

    def test_missing_project_raises(self, seeded):
        conn, root = seeded
        with pytest.raises(ProjectNotFound):
            service(conn, root).create_export_job(
                "prj_missing", ExportRequest(version=1, format="pptx")
            )

    def test_deadline_uses_rendering_budget(self, seeded):
        # docs/08 运行隔离：渲染时限 120 秒，不套用模型任务的 600 秒预算。
        conn, root = seeded
        job = service(conn, root).create_export_job(
            "prj_001", ExportRequest(version=1, format="pptx")
        )
        row = conn.execute(
            "SELECT created_at, deadline_at FROM jobs WHERE id=?", (job.id,)
        ).fetchone()
        created = datetime.fromisoformat(row["created_at"])
        deadline = datetime.fromisoformat(row["deadline_at"])
        delta = (deadline - created).total_seconds()
        assert delta == pytest.approx(EXPORT_DEADLINE_SECONDS, abs=2)


class TestHandleExport:
    def _run(self, conn, root):
        svc = service(conn, root)
        job = svc.create_export_job("prj_001", ExportRequest(version=1, format="pptx"))
        claimed = conn.execute("SELECT * FROM jobs WHERE id=?", (job.id,)).fetchone()
        return svc, job, claimed

    def test_produces_pptx_and_report_artifacts(self, seeded):
        conn, root = seeded
        svc, job, _ = self._run(conn, root)
        artifact_id = run_export(svc, conn, job.id)
        rows = conn.execute(
            "SELECT id, mime, source_type, version FROM artifacts WHERE project_id='prj_001'"
        ).fetchall()
        mimes = {r["mime"] for r in rows}
        assert "application/vnd.openxmlformats-officedocument.presentationml.presentation" in mimes
        assert "application/json" in mimes
        assert all(r["version"] == 1 for r in rows)
        # 返回的 artifact 可读且结构有效（L1 链完整由循环①测试锁定，这里验交付面）。
        store = ArtifactStore(conn, root)
        data = store.read_verified(artifact_id)
        assert data[:2] == b"PK"
        pres = Presentation(io.BytesIO(data))
        assert len(pres.slides) == 1
        page_text = "\n".join(
            sh.text_frame.text for sh in pres.slides[0].shapes if sh.has_text_frame
        )
        assert "页 one" in page_text

    def test_version_pinned_after_newer_commit(self, seeded):
        # 固定 version 语义：v2 已存在时导出 v1，内容必须是 v1。
        conn, root = seeded
        svc = service(conn, root)
        job = svc.create_export_job("prj_001", ExportRequest(version=1, format="pptx"))
        artifact_id = run_export(svc, conn, job.id)
        data = ArtifactStore(conn, root).read_verified(artifact_id)
        pres = Presentation(io.BytesIO(data))
        texts = "\n".join(
            sh.text_frame.text for sh in pres.slides[0].shapes if sh.has_text_frame
        )
        assert "页 one" in texts and "页 two" not in texts

    def test_evidence_report_content(self, seeded):
        conn, root = seeded
        svc = service(conn, root)
        job = svc.create_export_job("prj_001", ExportRequest(version=1, format="pptx"))
        run_export(svc, conn, job.id)
        row = conn.execute(
            "SELECT id FROM artifacts WHERE project_id='prj_001' AND mime='application/json'"
        ).fetchone()
        import json

        report = json.loads(ArtifactStore(conn, root).read_verified(row["id"]))
        assert report["version"] == 1
        assert report["exporter_version"] == EXPORTER_VERSION
        assert report["claims"] == []

    def test_corrupt_deck_fails_typed_without_artifact(self, seeded):
        # 损坏 committed 版本（ghost fact 引用）：导出显式失败，绝不产出残缺文件。
        conn, root = seeded
        corrupt = copy.deepcopy(deck(2, "坏"))
        payload = corrupt.model_dump(mode="json")
        payload["slides"][0]["blocks"] = [{"type": "fact", "claim_id": "ghost"}]
        with conn:
            conn.execute(
                "UPDATE deck_versions SET deck_json=? WHERE project_id='prj_001' AND version=2",
                (json.dumps(payload),),
            )
            assert conn.execute("SELECT changes() AS n").fetchone()["n"] == 1
        svc = service(conn, root)
        job = svc.create_export_job("prj_001", ExportRequest(version=2, format="pptx"))
        with pytest.raises(DeckExportError):
            svc.handle_export(job.id)
        assert (
            conn.execute(
                "SELECT COUNT(*) AS n FROM artifacts WHERE project_id='prj_001'"
            ).fetchone()["n"]
            == 0
        )

    def test_export_three_times_semantic_equal(self, seeded):
        # docs/08 L4：同输入多次导出，语义/结构一致（ZIP 时间戳差异允许）。
        conn, root = seeded
        svc = service(conn, root)
        structures = []
        for i in range(3):
            job = svc.create_export_job("prj_001", ExportRequest(version=1, format="pptx"))
            aid = run_export(svc, conn, job.id)
            data = ArtifactStore(conn, root).read_verified(aid)
            pres = Presentation(io.BytesIO(data))
            structures.append(
                [
                    [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
                    for s in pres.slides
                ]
            )
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                structures[-1].append(zf.read("docProps/core.xml"))
        assert structures[0] == structures[1] == structures[2]


class TestCorruptAndCancelGuards:
    def test_unparseable_deck_rejected_controlled_at_acceptance(self, seeded):
        # 结构层损坏（连 DeckSpec 校验都过不了）：受理期受控失败 EXPORT_FAILED，
        # 不得裸冒 ValidationError 变 HTTP 500 INTERNAL_ERROR（Review N1，
        # 与 previews 侧"诚实 failed"同口径）。
        conn, root = seeded
        with conn:
            conn.execute(
                "UPDATE deck_versions SET deck_json='{\"broken\": true}'"
                " WHERE project_id='prj_001' AND version=1",
            )
            assert conn.execute("SELECT changes() AS n").fetchone()["n"] == 1
        with pytest.raises(DeckExportError):
            service(conn, root).create_export_job(
                "prj_001", ExportRequest(version=1, format="pptx")
            )

    def test_cancelled_job_writes_no_artifacts(self, seeded):
        # 取消落在"渲染中→落盘"窗口：落盘前复核执行态，孤儿产物不产生（Review N3）。
        from courseware_core.errors import JobCancelled

        conn, root = seeded
        svc = service(conn, root)
        job = svc.create_export_job("prj_001", ExportRequest(version=1, format="pptx"))
        with conn:
            conn.execute(
                "UPDATE jobs SET status='running', worker_id='w_test', cancel_requested=1"
                " WHERE id=?",
                (job.id,),
            )
            assert conn.execute("SELECT changes() AS n").fetchone()["n"] == 1
        with pytest.raises(JobCancelled):
            svc.handle_export(job.id)
        assert (
            conn.execute("SELECT COUNT(*) AS n FROM artifacts").fetchone()["n"] == 0
        )

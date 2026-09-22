"""T10 循环①：python-pptx 导出器核心（docs/08 四类布局/确定性/预算/引用页脚）。

L1=ZIP/OOXML 完整部件链（T02 v1 缺链打不开的反例教训）；L2=文本对象可编辑
（读回断言 text_frame 而非图片）；L4=同输入语义/结构确定（docs/08 允许 ZIP
时间戳差异，文档属性时间戳必须固定）。
"""

import io
import zipfile

import pytest
from pptx import Presentation

from courseware_core.errors import DeckExportError
from courseware_core.models.common import CourseBrief
from courseware_core.models.deck import (
    DeckSpec,
    FactBlock,
    IllustrationBlock,
    Slide,
    TeachingBlock,
)
from courseware_core.models.evidence import Claim, EvidenceSpan
from courseware_core.render.exporter import (
    EXPORTER_VERSION,
    build_evidence_report,
    document_codes,
    export_deck_pptx,
)


def span(doc="mat_a", page=3, chunk="ck_1", quote="优先级分组"):
    return EvidenceSpan(
        chunk_id=chunk, document_id=doc, pdf_page=page, start=0, end=4, quote=quote
    )


def claim(cid="c1", text="NVIC_SetPriority 设置抢占优先级。", refs=None):
    return Claim(
        id=cid,
        text=text,
        kind="direct",
        evidence_refs=refs if refs is not None else [span()],
        rationale=None,
    )


def make_deck(slides=None, claims=None):
    return DeckSpec(
        schema_version="1.0.0",
        project_id="prj_1",
        version=2,
        corpus_revision=1,
        course=CourseBrief(
            topic="STM32 GPIO 与中断",
            audience="嵌入式系统大二学生",
            duration_minutes=45,
            goals=["掌握 GPIO 八种模式", "理解中断分组"],
            target_slides=4,
        ),
        claims=claims if claims is not None else [claim()],
        slides=slides
        if slides is not None
        else [
            Slide(
                id="s1",
                title="课程导入",
                layout="title",
                blocks=[
                    TeachingBlock(type="teaching", text="本讲从中断寄存器切入，45 分钟。")
                ],
            )
        ],
    )


def open_deck(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


class TestL1PackageStructure:
    def test_full_ooxml_chain_present(self):
        data = export_deck_pptx(make_deck())
        assert data[:2] == b"PK"
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = set(zf.namelist())
        # T02 v1 反例教训：缺 slideMaster/theme 链的 PPTX 打不开，必须全链在位。
        for required in (
            "[Content_Types].xml",
            "_rels/.rels",
            "ppt/presentation.xml",
            "ppt/slideMasters/slideMaster1.xml",
            "ppt/slideLayouts/slideLayout1.xml",
            "ppt/theme/theme1.xml",
            "ppt/slides/slide1.xml",
        ):
            assert required in names, required

    def test_16_9_geometry(self):
        pres = open_deck(export_deck_pptx(make_deck()))
        from pptx.util import Inches

        assert pres.slide_width == Inches(13.333) or pres.slide_width >= Inches(13.3)
        assert pres.slide_height == Inches(7.5)


class TestFourLayouts:
    def _slides(self):
        return [
            Slide(
                id="t1",
                title="封面",
                layout="title",
                blocks=[TeachingBlock(type="teaching", text="教学说明一段。")],
            ),
            Slide(
                id="c1",
                title="关键概念",
                layout="concept",
                blocks=[
                    FactBlock(type="fact", claim_id="c1"),
                    TeachingBlock(type="teaching", text="补充解释。"),
                ],
            ),
            Slide(
                id="w1",
                title="对比",
                layout="two_column",
                blocks=[
                    TeachingBlock(type="teaching", text="左组内容。"),
                    TeachingBlock(type="teaching", text="右组内容。"),
                ],
            ),
            Slide(
                id="p1",
                title="流程与示例",
                layout="process_example",
                blocks=[
                    TeachingBlock(type="teaching", text="第一步配置时钟。"),
                    IllustrationBlock(
                        type="illustration",
                        text="课堂问题：为什么先分组？",
                        assumptions=["假设板卡已上电"],
                        evidence_refs=[span()],
                    ),
                ],
            ),
        ]

    def test_each_layout_renders_with_title(self):
        data = export_deck_pptx(make_deck(slides=self._slides()))
        pres = open_deck(data)
        assert len(pres.slides) == 4
        texts = [
            "\n".join(
                sh.text_frame.text for sh in s.shapes if sh.has_text_frame
            )
            for s in pres.slides
        ]
        for slide, page_text in zip(self._slides(), texts):
            assert slide.title in page_text
        # fact 块展开为 claim 文本（claim 集合保留，非静默丢内容）。
        assert "NVIC_SetPriority" in texts[1]
        # illustration 假设进入页面（不伪造、显式标注）。
        assert "假设板卡已上电" in texts[3]


class TestChineseLongAndMultiline:
    def test_preserved_verbatim(self):
        long_title = "中断控制器与优先级分组机制详解及课堂演练环节安排说明补充" + "字" * 12
        multiline = (
            "第一行包含英文 API 名 NVIC_SetPriority 与扩展汉字𠀀𠮷。\n"
            "第二行中文换行测试，覆盖 200 字符上限附近的长教学文本。" + "补充" * 60
        )[:200]
        slides = [
            Slide(
                id="s1",
                title=long_title,
                layout="concept",
                blocks=[TeachingBlock(type="teaching", text=multiline)],
            )
        ]
        pres = open_deck(export_deck_pptx(make_deck(slides=slides)))
        all_text = "\n".join(
            sh.text_frame.text for sh in pres.slides[0].shapes if sh.has_text_frame
        )
        assert long_title in all_text
        assert multiline in all_text


class TestFooterCitation:
    def test_document_codes_are_deterministic(self):
        deck = make_deck(
            slides=[
                Slide(
                    id="s1",
                    title="概念",
                    layout="concept",
                    blocks=[FactBlock(type="fact", claim_id="c1")],
                )
            ],
            claims=[
                claim("c1", refs=[span(doc="mat_b"), span(doc="mat_a", page=7)]),
            ],
        )
        codes = document_codes(deck)
        # 稳定分配：按 document_id 排序编号，与出现顺序无关。
        assert codes == {"mat_a": "M1", "mat_b": "M2"}
        pres = open_deck(export_deck_pptx(deck))
        page_text = "\n".join(
            sh.text_frame.text for sh in pres.slides[0].shapes if sh.has_text_frame
        )
        assert "M1·p7" in page_text and "M2·p3" in page_text


class TestDeterminism:
    def _semantic_structure(self, data: bytes):
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            core = zf.read("docProps/core.xml")
        pres = open_deck(data)
        return (
            len(pres.slides),
            [
                [
                    (sh.shape_id, sh.text_frame.text if sh.has_text_frame else None)
                    for sh in s.shapes
                ]
                for s in pres.slides
            ],
            core,
        )

    def test_same_deck_same_semantic_structure_and_doc_props(self):
        deck = make_deck()
        a = export_deck_pptx(deck)
        b = export_deck_pptx(deck)
        # docs/08 L4：允许 ZIP 时间戳差异，比较语义/结构；文档属性时间戳
        # 属语义输出的一部分，必须固定（不随导出时刻漂移）。
        assert self._semantic_structure(a) == self._semantic_structure(b)


class TestTextSafety:
    def test_xml_special_text_not_injected(self):
        tricky = "<script>alert(1)</script> 与 & 符号"
        slides = [
            Slide(
                id="s1",
                title=tricky,
                layout="title",
                blocks=[TeachingBlock(type="teaching", text="正常正文")],
            )
        ]
        data = export_deck_pptx(make_deck(slides=slides))
        pres = open_deck(data)
        texts = [
            sh.text_frame.text for sh in pres.slides[0].shapes if sh.has_text_frame
        ]
        assert any(tricky in t for t in texts)
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            slide_xml = zf.read("ppt/slides/slide1.xml").decode("utf-8")
        assert "<script>alert" not in slide_xml

    def test_editable_text_frames_not_images(self):
        data = export_deck_pptx(make_deck())
        pres = open_deck(data)
        for slide in pres.slides:
            picture_shapes = [sh for sh in slide.shapes if sh.shape_type == 13]
            assert not picture_shapes
            text_frames = [sh for sh in slide.shapes if sh.has_text_frame]
            assert any(sh.text_frame.text.strip() for sh in text_frames)


class TestTheme:
    def test_exporter_version_and_theme_declared(self):
        assert EXPORTER_VERSION.startswith("python-pptx-")
        assert "classroom_v1" in EXPORTER_VERSION

    def test_title_uses_dark_body_blue_accent(self):
        from pptx.util import Pt

        deck = make_deck(
            slides=[
                Slide(
                    id="s1",
                    title="标题",
                    layout="concept",
                    blocks=[FactBlock(type="fact", claim_id="c1")],
                )
            ]
        )
        pres = open_deck(export_deck_pptx(deck))
        runs = [
            run
            for sh in pres.slides[0].shapes
            if sh.has_text_frame
            for p in sh.text_frame.paragraphs
            for run in p.runs
        ]
        colors = [run.font.color.rgb for run in runs if run.font.color.rgb is not None]
        assert colors, "导出必须显式声明字体颜色（classroom_v1）"
        from courseware_core.render.exporter import TEXT_DARK

        assert TEXT_DARK in colors


class TestBudgetBoundary:
    def test_max_deck_16_slides_6_blocks(self):
        claims = [claim(f"c{i}", text="事" * 160) for i in range(96)]
        slides = [
            Slide(
                id=f"s{k}",
                title="标" * 40,
                layout="two_column",
                blocks=[FactBlock(type="fact", claim_id=f"c{k*6+j}") for j in range(6)],
            )
            for k in range(16)
        ]
        deck = make_deck(slides=slides, claims=claims)
        pres = open_deck(export_deck_pptx(deck))
        assert len(pres.slides) == 16


class TestStructuralGuards:
    def test_unresolvable_fact_claim_raises_typed_error(self):
        slides = [
            Slide(
                id="s1",
                title="坏引用",
                layout="concept",
                blocks=[FactBlock(type="fact", claim_id="ghost")],
            )
        ]
        with pytest.raises(DeckExportError):
            export_deck_pptx(make_deck(slides=slides, claims=[claim("c1")]))

    def test_empty_claims_title_only_deck_still_exports(self):
        # title 页可零 fact 零引用（封面合法），不得误伤。
        pres = open_deck(export_deck_pptx(make_deck(claims=[])))
        assert len(pres.slides) == 1


class TestEvidenceReport:
    def test_report_covers_all_claims_with_codes(self):
        deck = make_deck(
            slides=[
                Slide(
                    id="s1",
                    title="概念",
                    layout="concept",
                    blocks=[FactBlock(type="fact", claim_id="c1")],
                )
            ],
            claims=[claim("c1", refs=[span(doc="mat_a", page=7)])],
        )
        report = build_evidence_report(deck)
        assert report["version"] == deck.version
        assert report["document_codes"] == {"mat_a": "M1"}
        entry = report["claims"][0]
        assert entry["claim_id"] == "c1"
        assert entry["evidence"][0]["document_code"] == "M1"
        assert entry["evidence"][0]["pdf_page"] == 7
        assert entry["evidence"][0]["quote"] == "优先级分组"

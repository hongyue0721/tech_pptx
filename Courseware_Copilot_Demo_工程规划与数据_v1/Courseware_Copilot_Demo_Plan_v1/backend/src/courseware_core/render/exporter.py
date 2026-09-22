"""T10 导出器：DeckSpec → PPTX（ADR-11：python-pptx==1.0.2，唯一激活出口）。

docs/08 约束的落点：固定 16:9 与四类布局的坐标/字号/颜色全部来自本模块常量
（确定性布局函数，不随内容浮动）；正文 24pt、标题 32pt、页脚 12pt（正文硬
下限 22pt）；classroom_v1 主题=白底、深色正文、蓝色点缀；引用页脚只放资料
短代号+物理页，完整 quote 随 evidence-report.json 交付。

结构性防线（DeckExportError）：committed deck 若含不可解析的 fact 引用或重复
claim id（如 T12 前入库的损坏版本），导出显式失败——绝不静默丢内容或多义
引用任选一条，宁可不出，不出残缺。
"""

import io
from dataclasses import field
from datetime import datetime, timezone
from typing import Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pydantic import BaseModel, ConfigDict

from courseware_core.errors import DeckExportError
from courseware_core.models.deck import (
    DeckSpec,
    FactBlock,
    IllustrationBlock,
    Slide,
    TeachingBlock,
)
from courseware_core.models.evidence import EvidenceSpan

EXPORTER_VERSION = "python-pptx-1.0.2-classroom_v1-v1"

TEXT_DARK = RGBColor(0x1F, 0x29, 0x37)
ACCENT_BLUE = RGBColor(0x25, 0x63, 0xEB)
# 只声明字体名、不分发字体文件（THIRD_PARTY 字体条目）；Office 端缺字回退
# 由 L3 目标 Office 抽验覆盖，不在服务端伪造"保证渲染一致"。
FONT_NAME = "Microsoft YaHei"

TITLE_PT = 32
BODY_PT = 24
FOOTER_PT = 12

# 16:9 标准 EMU 尺寸与固定版面几何（侧边距 0.45 英寸，docs/08 §四类布局）。
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)
MARGIN = Inches(0.45)
CONTENT_W = SLIDE_W - 2 * MARGIN
BODY_TOP = Inches(1.55)
FOOTER_TOP = Inches(6.92)
FOOTER_H = Inches(0.42)

# 文档属性时间戳固定：导出是 committed version 的纯渲染，输出不得随导出
# 时刻漂移（docs/08 L4 语义确定口径；ZIP 容器时间戳差异由库决定、可接受）。
_DOC_EPOCH = datetime(1970, 1, 1, tzinfo=timezone.utc)

# evidence-report 的引用视图：只暴露交付需要的字段，不外泄偏移内部细节。
class _ReportEvidence(BaseModel):
    model_config = ConfigDict(frozen=True)

    document_code: str
    document_id: str
    pdf_page: int
    quote: str


class EvidenceReport(BaseModel):
    """伴随 evidence-report.json 的契约结构（docs/08：页脚短代号，完整
    quote 与核验材料在此，按 claim 全量导出、不省略）。"""

    model_config = ConfigDict(frozen=True)

    project_id: str
    version: int
    corpus_revision: int
    exporter_version: str
    document_codes: dict[str, str]
    claims: list[dict] = field(default_factory=list)


class _DeckContext:
    """一次导出的只读上下文：claim 文本/引用索引与资料代号。"""

    def __init__(self, deck: DeckSpec):
        self.texts: dict[str, str] = {}
        self.refs: dict[str, list[EvidenceSpan]] = {}
        for claim in deck.claims:
            if claim.id in self.texts:
                raise DeckExportError(
                    "deck contains duplicate claim ids",
                    {"claim_id": claim.id, "version": deck.version},
                )
            self.texts[claim.id] = claim.text
            self.refs[claim.id] = list(claim.evidence_refs)
        self.codes = {
            doc_id: f"M{i}"
            for i, doc_id in enumerate(sorted(self._document_ids(deck)), 1)
        }

    @staticmethod
    def _document_ids(deck: DeckSpec):
        docs = {ref.document_id for claim in deck.claims for ref in claim.evidence_refs}
        for slide in deck.slides:
            for block in slide.blocks:
                if isinstance(block, IllustrationBlock):
                    docs.update(ref.document_id for ref in block.evidence_refs)
        return docs

    def claim_text(self, claim_id: str) -> str:
        if claim_id not in self.texts:
            # 引用不可解析=损坏 deck：显式失败而非渲染空白事实。
            raise DeckExportError(
                "slide references a claim id that does not exist in deck",
                {"claim_id": claim_id},
            )
        return self.texts[claim_id]

    def claim_refs(self, claim_id: str) -> list[EvidenceSpan]:
        return self.refs.get(claim_id, [])


def document_codes(deck: DeckSpec) -> dict[str, str]:
    """资料短代号：按 document_id 排序稳定分配 M1..Mn（与出现顺序无关）。"""
    return _DeckContext(deck).codes


def _text_items(slide: Slide, ctx: _DeckContext) -> list[tuple[str, int, bool]]:
    """块序列 → (文本, 字号, 加粗)。fact 展开为 claim 正文（claim 集合保留，
    不静默丢内容）；illustration 附课堂假设标注（显式假设，不伪造）；
    process_example 的步骤块带序号。"""
    items: list[tuple[str, int, bool]] = []
    step_no = 0
    for block in slide.blocks:
        if isinstance(block, FactBlock):
            text = ctx.claim_text(block.claim_id)
            if slide.layout == "process_example":
                step_no += 1
                text = f"{step_no}. {text}"
            items.append((text, BODY_PT, False))
        elif isinstance(block, TeachingBlock):
            text = block.text
            if slide.layout == "process_example":
                step_no += 1
                text = f"{step_no}. {text}"
            items.append((text, BODY_PT, False))
        elif isinstance(block, IllustrationBlock):
            items.append((block.text, BODY_PT, False))
            for assumption in block.assumptions:
                items.append((f"课堂假设：{assumption}", BODY_PT, False))
    return items


def _slide_ref_pairs(slide: Slide, ctx: _DeckContext) -> list[tuple[str, int]]:
    """页脚引用集：本页 fact 所属 claim 的 refs + illustration 直挂 refs，
    按 (document_id, pdf_page) 去重排序稳定输出。"""
    pairs: set[tuple[str, int]] = set()
    for block in slide.blocks:
        if isinstance(block, FactBlock):
            refs = ctx.claim_refs(block.claim_id)
        elif isinstance(block, IllustrationBlock):
            refs = block.evidence_refs
        else:
            continue
        pairs.update((ref.document_id, ref.pdf_page) for ref in refs)
    return sorted(pairs)


def _set_east_asian_font(run) -> None:
    # python-pptx 的 font.name 只写 a:latin 型face；中文渲染取 a:ea 型face，
    # 不补则 Office 端中文回退默认字体，classroom_v1 一致性无保证。
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT_NAME)


def _write_lines(tf, lines: list[tuple[str, int, bool]], align: Optional[PP_ALIGN] = None) -> None:
    for i, (text, size_pt, bold) in enumerate(lines):
        # 手动按 \n 拆段：读回 text 与输入逐行一致，不依赖库对控制符的处理。
        chunks = text.split("\n")
        for j, chunk in enumerate(chunks):
            para = tf.paragraphs[0] if (i == 0 and j == 0) else tf.add_paragraph()
            if align is not None:
                para.alignment = align
            run = para.add_run()
            run.text = chunk
            run.font.size = Pt(size_pt)
            run.font.bold = bold
            run.font.name = FONT_NAME
            run.font.color.rgb = TEXT_DARK
            _set_east_asian_font(run)


def _add_box(page, left, top, width, height):
    box = page.shapes.add_textbox(left, top, width, height)
    box.text_frame.word_wrap = True
    return box


def _add_accent_bar(page, left, top, width):
    bar = page.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_BLUE
    bar.line.fill.background()


def _add_footer(page, ctx: _DeckContext, slide: Slide) -> None:
    refs = _slide_ref_pairs(slide, ctx)
    if not refs:
        return
    label = "; ".join(f"{ctx.codes[doc]}·p{page_no}" for doc, page_no in refs)
    box = _add_box(page, MARGIN, FOOTER_TOP, CONTENT_W, FOOTER_H)
    _write_lines(box.text_frame, [(f"依据：{label}", FOOTER_PT, False)])


def _title_box(page, s: Slide, cover: bool = False):
    if cover:
        box = _add_box(page, MARGIN, Inches(2.5), CONTENT_W, Inches(1.4))
        _write_lines(box.text_frame, [(s.title, TITLE_PT, True)], align=PP_ALIGN.CENTER)
        _add_accent_bar(page, Inches(5.92), Inches(3.95), Inches(1.5))
    else:
        box = _add_box(page, MARGIN, Inches(0.4), CONTENT_W, Inches(0.95))
        _write_lines(box.text_frame, [(s.title, TITLE_PT, True)])
        _add_accent_bar(page, MARGIN, Inches(1.38), CONTENT_W)


def _render_title_layout(page, s: Slide, deck: DeckSpec, ctx: _DeckContext) -> None:
    # 封面：大标题+教学说明居中，副题带课程对象与时长（来自 course 真值）。
    _title_box(page, s, cover=True)
    subtitle = f"{deck.course.audience} · {deck.course.duration_minutes} 分钟"
    box = _add_box(page, MARGIN, Inches(4.35), CONTENT_W, Inches(1.8))
    _write_lines(box.text_frame, [(subtitle, BODY_PT, False)] + _text_items(s, ctx), align=PP_ALIGN.CENTER)
    _add_footer(page, ctx, s)


def _render_bulleted(page, s: Slide, ctx: _DeckContext) -> None:
    _title_box(page, s)
    box = _add_box(page, MARGIN, BODY_TOP, CONTENT_W, FOOTER_TOP - BODY_TOP - Inches(0.1))
    _write_lines(box.text_frame, _text_items(s, ctx))
    _add_footer(page, ctx, s)


def _render_two_column(page, s: Slide, ctx: _DeckContext) -> None:
    _title_box(page, s)
    items = _text_items(s, ctx)
    half = (len(items) + 1) // 2
    col_w = (CONTENT_W - Inches(0.3)) // 2
    body_h = FOOTER_TOP - BODY_TOP - Inches(0.1)
    left = _add_box(page, MARGIN, BODY_TOP, col_w, body_h)
    _write_lines(left.text_frame, items[:half])
    right = _add_box(page, MARGIN + col_w + Inches(0.3), BODY_TOP, col_w, body_h)
    _write_lines(right.text_frame, items[half:])
    _add_footer(page, ctx, s)


def export_deck_pptx(deck: DeckSpec) -> bytes:
    ctx = _DeckContext(deck)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for s in deck.slides:
        page = prs.slides.add_slide(blank)
        if s.layout == "title":
            _render_title_layout(page, s, deck, ctx)
        elif s.layout == "two_column":
            _render_two_column(page, s, ctx)
        else:
            _render_bulleted(page, s, ctx)

    props = prs.core_properties
    props.title = deck.course.topic
    props.author = "Courseware Copilot"
    props.subject = f"deck v{deck.version} corpus r{deck.corpus_revision}"
    props.comments = f"exporter={EXPORTER_VERSION}"
    props.created = _DOC_EPOCH
    props.modified = _DOC_EPOCH
    props.revision = deck.version

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_evidence_report(deck: DeckSpec) -> dict:
    ctx = _DeckContext(deck)
    return EvidenceReport(
        project_id=deck.project_id,
        version=deck.version,
        corpus_revision=deck.corpus_revision,
        exporter_version=EXPORTER_VERSION,
        document_codes=ctx.codes,
        claims=[
            {
                "claim_id": claim.id,
                "text": claim.text,
                "kind": claim.kind,
                "evidence": [
                    _ReportEvidence(
                        document_code=ctx.codes[ref.document_id],
                        document_id=ref.document_id,
                        pdf_page=ref.pdf_page,
                        quote=ref.quote,
                    ).model_dump()
                    for ref in claim.evidence_refs
                ],
            }
            for claim in deck.claims
        ],
    ).model_dump()

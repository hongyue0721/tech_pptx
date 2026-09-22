"""T10 导出服务：固定 version 的 PPTX 渲染 job（api.md §导出）。

export 是只读快照 job：受理不占项目写锁（claim_next 对 kind='export' 放行，
排队不排他）；渲染输入是 deck_versions 的 committed 内容，不重新推理、不调
模型、不外发内容（无 consent 门）。产物经 ArtifactStore 原子落盘：pptx 与
evidence-report.json 成对交付，job.result_ref 指向 pptx artifact。

L1 结构自检（validation_status=passed）只声明"ZIP/OOXML 部件链与页数自洽"，
不冒充 L3 目标 Office 人工验收（docs/08 验收层次）。
"""

import io
import json
import secrets
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from pptx import Presentation

from courseware_core.errors import (
    DeckExportError,
    DeckNotFound,
    JobCancelled,
    JobNotFound,
    ProjectNotFound,
)
from courseware_core.models import DeckSpec, Job
from courseware_core.models.export import ExportRequest
from courseware_core.render.exporter import EXPORTER_VERSION, build_evidence_report, export_deck_pptx
from courseware_core.storage.artifact_store import ArtifactStore
from courseware_core.storage.job_repository import JobRepository
from courseware_core.storage.project_repository import ProjectRepository
from courseware_core.storage.version_repository import VersionRepository

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# docs/08 运行隔离：渲染任务时限 120 秒（M0 校准值），与模型任务 600 秒预算分开。
EXPORT_DEADLINE_SECONDS = 120


class ExportService:
    def __init__(self, conn, artifacts_root: Path):
        self._conn = conn
        self._projects = ProjectRepository(conn)
        self._versions = VersionRepository(conn)
        # 渲染无模型调用：deadline 用 120s 预算构造（JobRepository 按 created_at 推算）。
        self._jobs = JobRepository(conn, job_deadline_seconds=EXPORT_DEADLINE_SECONDS)
        self._artifacts_root = Path(artifacts_root)

    def _load_deck(self, project_id: str, version: int) -> DeckSpec:
        # 结构层损坏（deck_json 连模型校验都过不了）归一为受控失败：受理期
        # 显式 500/EXPORT_FAILED（数据完整性族，error_mapping 登记，拒绝隐式
        # 兜底与无信息的 INTERNAL_ERROR）、执行期 job=failed/EXPORT_FAILED
        #（Review N1，与 previews 侧"诚实 failed"同口径）。
        try:
            deck = self._versions.get_deck(project_id, version)
        except DeckExportError:
            raise
        except Exception as exc:
            raise DeckExportError(
                "stored deck could not be parsed",
                {"project_id": project_id, "version": version, "cause": type(exc).__name__},
            ) from exc
        if deck is None:
            raise DeckNotFound({"project_id": project_id, "version": version})
        return deck

    def create_export_job(
        self, project_id: str, request: ExportRequest, request_id: Optional[str] = None
    ) -> Job:
        project = self._projects.get(project_id)
        if project is None:
            raise ProjectNotFound({"project_id": project_id})
        deck = self._load_deck(project_id, request.version)
        now = datetime.now(timezone.utc)
        job = Job(
            id=f"job_{secrets.token_hex(16)}",
            project_id=project_id,
            kind="export",
            status="queued",
            stage="queued",
            cancel_requested=False,
            base_version=request.version,
            # 固定 version 语义：corpus 取该版本入库时点，而非项目当前值。
            corpus_revision=deck.corpus_revision,
            created_at=now,
            updated_at=now,
        )
        self._jobs.create(job, request_id=request_id)
        return job

    def handle_export(self, job_id: str) -> str:
        """worker handler 体：渲染 + L1 自检 + 双 artifact 原子落盘。返回 pptx artifact_id。"""
        job = self._jobs.get(job_id)
        if job is None:
            raise JobNotFound({"job_id": job_id})
        self._jobs.set_stage(job_id, "rendering")
        deck = self._load_deck(job.project_id, job.base_version)

        data = export_deck_pptx(deck)
        self._l1_self_check(data, deck)

        # 落盘前复核执行态：取消/中断落在"渲染→写库"窗口时不产孤儿产物
        #（Review N3；worker _publish 终判仍在，这里只是提前止损）。
        state = self._jobs.get_execution_state(job_id)
        if state is None or state["status"] != "running" or state["cancel_requested"]:
            raise JobCancelled({"job_id": job_id})

        suffix = job.id.removeprefix("job_")
        pptx_id = f"art_{suffix}_pptx"
        report_id = f"art_{suffix}_report"
        store = ArtifactStore(self._conn, self._artifacts_root)
        store.put(
            artifact_id=pptx_id,
            project_id=job.project_id,
            version=job.base_version,
            corpus_revision=job.corpus_revision,
            data=data,
            mime=PPTX_MIME,
            source_type="export",
            validation_status="passed",
            renderer_version=EXPORTER_VERSION,
            font_profile="declared:Microsoft YaHei;embedded:false",
        )
        store.put(
            artifact_id=report_id,
            project_id=job.project_id,
            version=job.base_version,
            corpus_revision=job.corpus_revision,
            data=json.dumps(
                build_evidence_report(deck), ensure_ascii=False, sort_keys=True
            ).encode("utf-8"),
            mime="application/json",
            source_type="export",
            validation_status="not_applicable",
            renderer_version=EXPORTER_VERSION,
        )
        return pptx_id

    @staticmethod
    def _l1_self_check(data: bytes, deck: DeckSpec) -> None:
        # 交付前结构自检：能重开、页数一致、部件链完整——坏文件宁可不交付。
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
                for required in (
                    "[Content_Types].xml",
                    "ppt/presentation.xml",
                    "ppt/slideMasters/slideMaster1.xml",
                    "ppt/theme/theme1.xml",
                ):
                    if required not in names:
                        raise DeckExportError(
                            "exported package failed L1 structure self-check",
                            {"missing_part": required},
                        )
            pres = Presentation(io.BytesIO(data))
        except DeckExportError:
            raise
        except Exception as exc:
            raise DeckExportError(
                "exported package failed L1 structure self-check", {"cause": str(exc)}
            ) from exc
        if len(pres.slides) != len(deck.slides):
            raise DeckExportError(
                "exported slide count does not match deck",
                {"expected": len(deck.slides), "actual": len(pres.slides)},
            )
